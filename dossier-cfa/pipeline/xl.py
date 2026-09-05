# -*- coding: utf-8 -*-
"""Styles neutres WE-FORM pour le classeur analytique."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TEAL   = '1A8A87'
TEAL_P = 'E8F2F1'
ENCRE  = '16201F'
ENCRE2 = '46524F'
ENCRE3 = '78837F'
FILET  = 'E2E6E4'
FEUIL2 = 'F7F8F7'
BRIQUE = 'A8412C'

FONT = 'Calibri'
NUM  = '# ##0.00;[Red]-# ##0.00;"·"'
NUM0 = '# ##0;[Red]-# ##0;"·"'
PCT  = '0.00"' + ' %"'
DATE = 'DD/MM/YYYY'

def _b(c='FFFFFF'): return Side(style='thin', color=c)

def title(ws, row, text, sub=None, width=8):
    ws.cell(row=row, column=1, value=text).font = Font(FONT, 16, bold=True, color=TEAL)
    ws.row_dimensions[row].height = 24
    if sub:
        ws.cell(row=row + 1, column=1, value=sub).font = Font(FONT, 9.5, color=ENCRE2)
        ws.row_dimensions[row + 1].height = 14
        return row + 3
    return row + 2

def header(ws, row, labels, widths=None, aligns=None):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = Font(FONT, 9, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=TEAL)
        c.alignment = Alignment(horizontal=(aligns[i - 1] if aligns else 'left'),
                                vertical='center', wrap_text=True)
        c.border = Border(bottom=_b(TEAL), top=_b(TEAL))
    ws.row_dimensions[row].height = 26
    if widths:
        for i, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(i)].width = w
    return row + 1

def row(ws, r, values, fmts=None, bold=False, fill=None, color=None, top_rule=False):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=i, value=v)
        f = (fmts[i - 1] if fmts and i - 1 < len(fmts) else None)
        c.font = Font(FONT, 9, bold=bold, color=color or ENCRE)
        if f == 'n':   c.number_format = NUM;  c.alignment = Alignment(horizontal='right')
        elif f == 'n0':c.number_format = NUM0; c.alignment = Alignment(horizontal='right')
        elif f == 'p': c.number_format = PCT;  c.alignment = Alignment(horizontal='right')
        elif f == 'd': c.number_format = DATE; c.alignment = Alignment(horizontal='center')
        elif f == 'c': c.alignment = Alignment(horizontal='center')
        elif f == 'w': c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        else:          c.alignment = Alignment(horizontal='left')
        if fill: c.fill = PatternFill('solid', fgColor=fill)
        c.border = Border(bottom=_b(FILET), top=_b(TEAL if top_rule else 'FFFFFF'))
    return r + 1

def total(ws, r, values, fmts=None, label_fill=TEAL_P):
    return row(ws, r, values, fmts, bold=True, fill=label_fill, color=TEAL, top_rule=True)

def note(ws, r, text, n=8):
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(FONT, 8.5, italic=True, color=ENCRE3)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
    ws.row_dimensions[r].height = max(14, 12 * (1 + len(text) // 150))
    return r + 2

def finish(ws, freeze=None, cols=None):
    ws.sheet_view.showGridLines = False
    if freeze: ws.freeze_panes = freeze
    if cols:
        for i, w in enumerate(cols, start=1): ws.column_dimensions[get_column_letter(i)].width = w
